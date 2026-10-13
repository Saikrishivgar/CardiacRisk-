import io
import json
import logging
import os
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
import joblib
import numpy as np
import pandas as pd
from pypdf import PdfReader
import shap
from fastapi import FastAPI, HTTPException, UploadFile, File, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from .clinical_extractor import parse_clinical_text
from .pdf_generator import generate_cardiac_pdf_report

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger("cardiac_risk_api")

ROOT = Path(__file__).resolve().parents[1]
MODEL_DIR = ROOT / "models"
METRICS_PATH = ROOT / "reports" / "metrics.json"

try:
    META = json.loads(METRICS_PATH.read_text())
    INPUT_COLUMNS: List[str] = META["input_columns"]
except Exception as e:
    logger.error(f"Failed to load metrics metadata: {e}")
    INPUT_COLUMNS = []

FORBIDDEN_TARGETS = {"Cath", "CAD", "LAD", "LCX", "RCA"}

# Model pipelines
try:
    MODELS = {
        "CAD": joblib.load(MODEL_DIR / "cath_model.joblib"),
        "LAD": joblib.load(MODEL_DIR / "lad_model.joblib"),
        "LCX": joblib.load(MODEL_DIR / "lcx_model.joblib"),
        "RCA": joblib.load(MODEL_DIR / "rca_model.joblib"),
    }
    logger.info(f"Successfully loaded {len(MODELS)} trained ML models: {list(MODELS.keys())}")
except Exception as e:
    logger.error(f"Failed to load one or more ML models: {e}")
    MODELS = {}

# Pre-cache TreeExplainers at startup for fast inference
try:
    EXPLAINERS = {
        target: shap.TreeExplainer(model_pipe.named_steps["model"])
        for target, model_pipe in MODELS.items()
    }
    logger.info(f"Successfully initialized {len(EXPLAINERS)} SHAP TreeExplainers: {list(EXPLAINERS.keys())}")
except Exception as e:
    logger.error(f"Failed to initialize SHAP TreeExplainers: {e}")
    EXPLAINERS = {}

SHAP_SCALES = {
    "CAD": "probability",
    "LAD": "probability",
    "LCX": "log_odds",
    "RCA": "probability",
}

# Categorical columns in training pipeline
CAT_COLS = [
    "Sex", "Obesity", "CRF", "CVA", "Airway disease", "Thyroid Disease", "CHF", "DLP",
    "Weak Peripheral Pulse", "Lung rales", "Systolic Murmur", "Diastolic Murmur",
    "Dyspnea", "Atypical", "Nonanginal", "Exertional CP", "LowTH Ang", "LVH",
    "Poor R Progression", "BBB", "VHD"
]

BINARY_NUMERIC_COLS = {
    "DM", "HTN", "Current Smoker", "EX-Smoker", "FH", "Edema",
    "Typical Chest Pain", "Q Wave", "St Elevation", "St Depression", "Tinversion"
}

YN_CATEGORICAL_COLS = {
    "Obesity", "CRF", "CVA", "Airway disease", "Thyroid Disease", "CHF", "DLP",
    "Weak Peripheral Pulse", "Lung rales", "Systolic Murmur", "Diastolic Murmur",
    "Dyspnea", "Atypical", "Nonanginal", "LowTH Ang", "LVH", "Poor R Progression"
}

ALLOWED_SEX = {"Male", "Female"}
ALLOWED_BBB = {"N", "LBBB", "RBBB"}
ALLOWED_VHD = {"N", "mild", "Moderate", "Severe"}
ALLOWED_FUNCTION_CLASS = {0, 1, 2, 3}
ALLOWED_REGION_RWMA = {0, 1, 2, 3, 4}


def clean_feature_name(raw: str) -> str:
    """Format transformed pipeline feature names into clear, human-readable labels."""
    if raw.startswith("num__"):
        return raw[5:]
    if raw.startswith("cat__"):
        rem = raw[5:]
        for c in CAT_COLS:
            if rem.startswith(c + "_"):
                val = rem[len(c) + 1:]
                return f"{c} = {val}"
        return rem
    return raw


def get_risk_level(prob: float) -> str:
    """
    Model risk visualization tier thresholds.
    Configured for prototype visualization; not a substitute for clinical diagnostics.
    """
    if prob < 0.45:
        return "LOW"
    elif prob < 0.75:
        return "MODERATE"
    else:
        return "HIGH"


def explain(target: str, model_pipe, explainer, row_df: pd.DataFrame, top_n: int = 5):
    prep = model_pipe.named_steps["prep"]
    transformed = prep.transform(row_df)
    names = prep.get_feature_names_out()

    values = explainer.shap_values(transformed)
    values = np.asarray(values)

    # RandomForest binary classifier returns [n_samples, n_features, 2]
    # XGBoost returns [n_samples, n_features]
    if values.ndim == 3:
        values = values[:, :, 1]
    vals = values[0]

    order = np.argsort(np.abs(vals))[::-1][:top_n]
    return [
        {
            "feature": clean_feature_name(str(names[i])),
            "shap_value": round(float(vals[i]), 4),
        }
        for i in order
    ]


class PredictionRequest(BaseModel):
    features: Dict[str, Any]


class GenerateReportRequest(BaseModel):
    features: Dict[str, Any]
    provided_features: Optional[List[str]] = None
    extraction_sources: Optional[Dict[str, str]] = None
    patient_identifier: Optional[str] = None
    doctor_user: Optional[str] = None
    vessel_image_base64: Optional[str] = None
    predictions: Optional[Dict[str, float]] = None
    risk_levels: Optional[Dict[str, str]] = None
    explanations: Optional[Dict[str, List[Dict[str, Any]]]] = None
    shap_scales: Optional[Dict[str, str]] = None


class ClinicalExtractionRequest(BaseModel):
    text: str


class LoginRequest(BaseModel):
    username: str
    password: str


app = FastAPI(title="Explainable 3D Cardiac Risk Navigator API")

# Configure CORS for production cloud deployment and local development
allowed_origins_env = os.getenv("ALLOWED_ORIGINS", "")
custom_origins = [orig.strip() for orig in allowed_origins_env.split(",") if orig.strip()]
default_origins = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
]
all_allowed_origins = list(set(default_origins + custom_origins))

app.add_middleware(
    CORSMiddleware,
    allow_origins=all_allowed_origins,
    allow_origin_regex=r"^https:\/\/.*\.vercel\.app$",
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


@app.get("/")
def root():
    return {
        "service": "Explainable 3D Cardiac Risk Navigator API",
        "status": "online",
        "health": "/health",
        "docs": "/docs",
        "version": "1.0.0",
    }


@app.get("/health")
def health():
    expected_models = {"CAD", "LAD", "LCX", "RCA"}
    models_status = {target: target in MODELS for target in expected_models}
    explainers_status = {target: target in EXPLAINERS for target in expected_models}

    all_ready = (
        all(models_status.values())
        and all(explainers_status.values())
        and len(INPUT_COLUMNS) > 0
    )

    if not all_ready:
        logger.error(f"Health check failed: models={models_status}, explainers={explainers_status}")
        raise HTTPException(
            status_code=503,
            detail={
                "status": "unhealthy",
                "message": "One or more machine learning models or explainers are not available.",
                "models_loaded": models_status,
                "explainers_loaded": explainers_status,
            },
        )

    return {
        "status": "ok",
        "service": "Explainable 3D Cardiac Risk Navigator API",
        "models_loaded": models_status,
        "explainers_loaded": explainers_status,
    }


@app.post("/auth/login")
def login(req: LoginRequest):
    user = req.username.strip()
    pwd = req.password
    if not user or not pwd:
        raise HTTPException(status_code=401, detail="Invalid username or password")
    if user == "saikrish" and pwd == "123456":
        return {
            "authenticated": True,
            "user": {
                "username": "saikrish",
                "name": "Dr. Saikrish (Cardiology)",
            },
        }
    raise HTTPException(status_code=401, detail="Invalid username or password")


def validate_and_run_inference(features: Dict[str, Any]) -> Dict[str, Any]:
    """Validates features against the 55-feature schema and executes authoritative model pipelines and SHAP explainers."""
    errors: List[str] = []

    # 1. Check for forbidden target columns
    found_targets = [col for col in features if col in FORBIDDEN_TARGETS]
    if found_targets:
        errors.append(f"Target columns must not be provided as inputs: {found_targets}")

    # 2. Check for unexpected / missing columns
    expected_set = set(INPUT_COLUMNS)
    missing = [c for c in INPUT_COLUMNS if c not in features]
    unexpected = [c for c in features if c not in expected_set]

    if missing:
        errors.append(f"Missing required features: {missing}")
    if unexpected:
        errors.append(f"Unexpected features provided: {unexpected}")

    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": "Feature validation failed", "errors": errors},
        )

    # 3. Strict type & value validations
    for col in INPUT_COLUMNS:
        val = features[col]

        # Numeric binary fields: must be numbers 0 or 1 (and not bool)
        if col in BINARY_NUMERIC_COLS:
            if isinstance(val, bool) or val not in (0, 1):
                errors.append(f"Field '{col}' must be numeric 0 or 1 (received: {repr(val)})")

        # Categorical Y/N fields: must be "Y" or "N"
        elif col in YN_CATEGORICAL_COLS:
            if val not in ("Y", "N"):
                errors.append(f"Field '{col}' must be 'Y' or 'N' (received: {repr(val)})")

        # Specific categorical fields
        elif col == "Sex":
            if val not in ALLOWED_SEX:
                errors.append(f"Field 'Sex' must be 'Male' or 'Female' (received: {repr(val)})")
        elif col == "BBB":
            if val not in ALLOWED_BBB:
                errors.append(f"Field 'BBB' must be one of {sorted(ALLOWED_BBB)} (received: {repr(val)})")
        elif col == "VHD":
            if val not in ALLOWED_VHD:
                errors.append(f"Field 'VHD' must be one of {sorted(ALLOWED_VHD)} (received: {repr(val)})")
        elif col == "Exertional CP":
            if val != "N":
                errors.append(f"Field 'Exertional CP' must be 'N' (received: {repr(val)})")
        elif col == "Function Class":
            if isinstance(val, bool) or val not in ALLOWED_FUNCTION_CLASS:
                errors.append(f"Field 'Function Class' must be an integer in 0–3 (received: {repr(val)})")
        elif col == "Region RWMA":
            if isinstance(val, bool) or val not in ALLOWED_REGION_RWMA:
                errors.append(f"Field 'Region RWMA' must be an integer in 0–4 (received: {repr(val)})")

        # Continuous numeric fields
        else:
            if isinstance(val, bool) or not isinstance(val, (int, float)):
                errors.append(f"Field '{col}' must be a numeric value (received: {repr(val)})")

    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": "Data type or category validation failed", "errors": errors},
        )

    # 4. Construct DataFrame in the exact training column order
    row = {c: features[c] for c in INPUT_COLUMNS}
    frame = pd.DataFrame([row])[INPUT_COLUMNS]

    predictions: Dict[str, float] = {}
    risk_levels: Dict[str, str] = {}
    explanations: Dict[str, List[Dict[str, Any]]] = {}

    for target, model_pipe in MODELS.items():
        prob = float(model_pipe.predict_proba(frame)[0, 1])
        predictions[target] = round(prob, 4)
        risk_levels[target] = get_risk_level(prob)
        explanations[target] = explain(target, model_pipe, EXPLAINERS[target], frame)

    return {
        "predictions": predictions,
        "risk_levels": risk_levels,
        "explanations": explanations,
        "shap_scales": SHAP_SCALES,
        "note": "Research/educational decision-support prototype; not a clinical diagnostic.",
    }


@app.post("/predict")
def predict(request: PredictionRequest):
    return validate_and_run_inference(request.features)


@app.post("/generate-report")
def generate_report(request: GenerateReportRequest):
    # Validate payload features and obtain authoritative model inference
    inference_result = validate_and_run_inference(request.features)

    # Determine unique report identifier
    report_id = request.patient_identifier or f"CRN-{datetime.now().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"
    filename = f"Cardiac_Risk_Report_{report_id}.pdf"

    # Assemble report generation payload
    report_payload = {
        "report_id": report_id,
        "doctor_user": request.doctor_user or "Dr. Saikrish (Cardiology)",
        "features": request.features,
        "provided_features": request.provided_features,
        "extraction_sources": request.extraction_sources,
        "predictions": inference_result["predictions"],
        "risk_levels": inference_result["risk_levels"],
        "explanations": inference_result["explanations"],
        "shap_scales": inference_result["shap_scales"],
        "vessel_image_base64": request.vessel_image_base64,
    }

    try:
        pdf_bytes = generate_cardiac_pdf_report(report_payload)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate clinical PDF report: {str(e)}")

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB

@app.post("/extract-text")
async def extract_text(file: UploadFile = File(...)):
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    
    if ext not in [".txt", ".pdf", ".docx"]:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}. Only .txt, .pdf, and .docx are supported.")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File size exceeds the 5MB limit.")
    if len(content) == 0:
        raise HTTPException(status_code=400, detail="File is empty.")

    extracted_text = ""
    try:
        if ext == ".txt":
            extracted_text = content.decode("utf-8")
        elif ext == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            extracted_text = "\n".join(pages_text)
        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.replace('\n', ' ').strip() for cell in row.cells])
                    if row_text.strip():
                        tables_text.append(row_text)
            extracted_text = "\n".join(paragraphs_text + tables_text)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Error extracting text from document: {str(e)}")
        
    return {"filename": filename, "extracted_text": extracted_text.strip()}


@app.post("/extract-clinical-data")
def extract_clinical_data(request: ClinicalExtractionRequest):
    if not request.text or not request.text.strip():
        raise HTTPException(status_code=400, detail="No text provided for extraction.")
        
    try:
        # Use our regex-based rule parser without LLM integration
        result = parse_clinical_text(request.text, INPUT_COLUMNS)
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to extract clinical data: {str(e)}")
